from pathlib import Path
import importlib,json,os,subprocess,sys,tempfile,unittest
from unittest.mock import patch
from PIL import Image,ImageDraw
from fig_cos.cli import skill_root

class Pipeline(unittest.TestCase):
    def command(self,*args,expected=0):
        result=subprocess.run([sys.executable,'-m','fig_cos',*map(str,args)],capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,expected,result.stdout+'\n'+result.stderr)
        return result

    def test_doctor(self):
        report=json.loads(self.command('doctor').stdout)
        self.assertTrue(report['core_ready']);self.assertEqual(report['default_reconstruction'],'Previous local tracing')
        self.assertFalse(report['powerpoint_application_verified'])

    def test_native_export_and_live_text(self):
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder)/'demo';self.command('demo','--output-dir',out)
            report=json.loads((out/'export/report.json').read_text(encoding='utf-8'))
            self.assertEqual(report['native_text_count'],6);self.assertEqual(report['raster_picture_count'],0)
            self.assertGreaterEqual(report['native_group_count'],2)
            from pptx import Presentation
            def texts(shapes):
                for shape in shapes:
                    if shape.shape_type==6:yield from texts(shape.shapes)
                    elif shape.has_text_frame:yield shape.text
            labels=list(texts(Presentation(str(out/'export/editable.pptx')).slides[0].shapes))
            self.assertIn('DNA helix',labels);self.assertIn('Cell cluster',labels)
            self.assertTrue((out/'svg-preview.png').exists())

    def test_original_resolution_provenance(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder);im=Image.new('RGB',(64,64),'white');d=ImageDraw.Draw(im)
            d.ellipse((8,8,54,54),fill='#65a7c1',outline='#285a77',width=2)
            raw=folder/'component.png';im.resize((128,128),Image.Resampling.LANCZOS).save(raw)
            out=folder/'trace';self.command('trace','--input-image',raw,'--text-free','--original-size',64,64,'--output-dir',out)
            report=json.loads((out/'provenance.json').read_text(encoding='utf-8'))
            self.assertEqual(report['original_pixel_dimensions'],[64,64]);self.assertTrue(report['input_resampled'])
            self.assertFalse(report['source_information_added_by_tracing']);self.assertGreater(report['path_count'],0)

    def test_installer_copies_resources_without_overwriting(self):
        with tempfile.TemporaryDirectory() as folder:
            target=Path(folder)/'fig-cos';self.command('install-skill','--target',target)
            self.assertIn('\nname: fig-cos\n',(target/'SKILL.md').read_text(encoding='utf-8'))
            self.assertTrue((target/'scripts/vendor/LICENSE').exists())
            marker=target/'keep.txt';marker.write_text('keep',encoding='utf-8')
            self.command('install-skill','--target',target,expected=2)
            self.assertEqual(marker.read_text(),'keep')

    def test_unedited_label_is_protected(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder);before=folder/'before.svg';after=folder/'after.svg'
            source='<svg xmlns="http://www.w3.org/2000/svg"><g id="selected"><path d="M0 0L2 2"/></g><text id="label">Fixed</text></svg>'
            before.write_text(source,encoding='utf-8');after.write_text(source.replace('L2 2','L4 4'),encoding='utf-8')
            self.command('check-preserved','--before',before,'--after',after,'--editable-component','selected')
            after.write_text(source.replace('Fixed','Changed'),encoding='utf-8')
            self.command('check-preserved','--before',before,'--after',after,'--editable-component','selected',expected=1)

    def test_offline_bundled_index_and_external_cache(self):
        sys.path.insert(0,str(skill_root()/'scripts'))
        library=importlib.import_module('icon_library')
        with tempfile.TemporaryDirectory() as folder,patch.object(library,'INDEX',Path(folder)/'absent.json'),patch.object(library,'get',side_effect=RuntimeError('Unexpected network')):
            self.assertTrue(library.search('antibody',2))
            self.assertFalse(library.CACHE.is_relative_to(skill_root()))

if __name__=='__main__':unittest.main()
