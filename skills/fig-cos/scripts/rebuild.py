"""Offline raster/SVG to editable SVG and PPTX. No network or credentials."""
from __future__ import annotations
import argparse
import copy
from collections import Counter
import hashlib
import importlib.metadata
import json
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
import native_gradients
from pathlib import Path

VENDOR = Path(__file__).resolve().parent / 'vendor'
ET.register_namespace('', 'http://www.w3.org/2000/svg')
PRESETS = {
    'balanced': dict(filter_speckle=2, color_precision=6, layer_difference=16,
                     length_threshold=4.0),
    'detail': dict(filter_speckle=1, color_precision=8, layer_difference=8,
                   length_threshold=2.0),
    'clean': dict(filter_speckle=8, color_precision=5, layer_difference=24,
                  length_threshold=5.0),
}


def run(script, *args):
    result = subprocess.run([sys.executable, '-X', 'utf8', str(VENDOR / script),
                             *map(str, args)], capture_output=True, text=True,
                            encoding='utf-8')
    if result.returncode:
        raise RuntimeError(f'{script} failed:\n{result.stderr}\n{result.stdout}')
    return result.stdout.strip()


def normalize_trace(path, width, height):
    tree = ET.parse(path)
    root = tree.getroot()
    root.set('viewBox', f'0 0 {width} {height}')
    for index, node in enumerate(root.iter()):
        if node.tag.rsplit('}', 1)[-1] == 'path' and not node.get('id'):
            node.set('id', f'trace-path-{index}')
    tree.write(path, encoding='utf-8', xml_declaration=True)


def close_implicit_fills(path):
    """SVG fills close open contours; preserve the separate open stroke."""
    cache=json.loads(path.read_text(encoding='utf-8'))
    atoms=[];indices={};changed=0
    for index,atom in enumerate(cache['atoms']):
        parts=atom.get('paintParts',[])
        open_fill=(atom['kind']=='path' and any(not p.get('closed') for p in atom.get('subpaths',[]))
                   and any(p.get('filled') for p in parts))
        indices[index]=[]
        def append(value):
            indices[index].append(len(atoms));atoms.append(value)
        if not open_fill:append(atom);continue
        changed+=1
        for paint in parts:
            if paint.get('filled'):
                filled=copy.deepcopy(atom);part=copy.deepcopy(paint);part['stroked']=False
                filled['paintParts']=[part]
                for contour in filled['subpaths']:contour['closed']=True
                append(filled)
            if paint.get('stroked'):
                stroked=copy.deepcopy(atom);part=copy.deepcopy(paint);part['filled']=False
                stroked['paintParts']=[part];append(stroked)
    if changed:
        cache['atoms']=atoms
        for batch in cache['batches']:
            batch['atom_indices']=[new for old in batch['atom_indices'] for new in indices[int(old)]]
        cache['implicit_fill_closed_atoms']=changed
        path.write_text(json.dumps(cache,ensure_ascii=False),encoding='utf-8')
    return changed


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument('--input-image', type=Path)
    source.add_argument('--input-svg', type=Path)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--preset', choices=PRESETS, default='detail')
    parser.add_argument('--text-manifest', type=Path)
    parser.add_argument('--text-already-removed', action='store_true')
    parser.add_argument('--max-paths', type=int, default=15000)
    args = parser.parse_args()
    input_path = (args.input_image or args.input_svg).resolve(strict=True)
    if args.text_manifest and args.input_image and not args.text_already_removed:
        parser.error('Use a verified text-free image and --text-already-removed with a manifest.')
    manifest = args.text_manifest.resolve(strict=True) if args.text_manifest else None
    if args.max_paths < 1:
        parser.error('--max-paths must be positive')
    output = args.output_dir.resolve()
    if output.exists():
        parser.error('Output directory already exists; select a fresh directory.')
    source_hash = hashlib.sha256(input_path.read_bytes()).hexdigest()
    output.mkdir(parents=True, exist_ok=False)
    shutil.copy2(input_path, output / ('source' + input_path.suffix.lower()))
    raw, master = output / 'traced.svg', output / 'editable.svg'
    warnings = []
    if args.input_image:
        import vtracer
        from PIL import Image
        with Image.open(input_path) as image:
            width, height = image.size
            if getattr(image, 'n_frames', 1) != 1:
                raise ValueError('Select a single image frame before tracing.')
        vtracer.convert_image_to_svg_py(str(input_path), str(raw),
            colormode='color', hierarchical='stacked', mode='spline',
            corner_threshold=60, max_iterations=10, splice_threshold=45,
            path_precision=3, **PRESETS[args.preset])
        normalize_trace(raw, width, height)
        if not manifest:
            warnings.append('No text manifest: any source lettering remains outlines, not editable typing.')
    else:
        shutil.copy2(input_path, raw)
    if manifest:
        copy = output / 'text-manifest.json'
        shutil.copy2(manifest, copy)
        run('merge_live_text.py', '--input-svg', raw, '--text-manifest', copy,
            '--output-svg', master)
    else:
        shutil.copy2(raw, master)
    validation = json.loads(run('validate_vector_svg.py', '--svg', master,
        '--report', output / 'svg-validation.json'))
    if validation['counts'].get('path', 0) > args.max_paths:
        raise ValueError('Too many paths for practical native PPT editing; try clean preset or rebuild key regions.')
    geometry_input,gradient_bindings=native_gradients.prepare(master,output/'geometry-input.svg')
    run('prepare_geometry_cache.py', '--input', geometry_input,
        '--output-dir', output / 'geometry', '--job-id', 'fig_cos')
    fill_closure_count=close_implicit_fills(output/'geometry'/'geometry-cache.json')
    pptx = output / 'editable.pptx'
    export = json.loads(run('run_cell_ppt_ooxml.py', '--geometry-cache',
        output / 'geometry' / 'geometry-cache.json', '--output-pptx', pptx))
    from pptx import Presentation
    deck = Presentation(str(pptx))
    slide = deck.slides[0]
    # SVG library instances become native PPT groups, so an icon can be moved
    # as one object while its individual vector paths remain editable.
    cache=json.loads((output/'geometry'/'geometry-cache.json').read_text(encoding='utf-8'))
    native_gradient_count=native_gradients.apply(slide,cache,gradient_bindings)
    source_atoms={atom['sourceId']:atom['objectName'] for atom in cache['atoms']}
    shapes_by_atom={}
    for shape in slide.shapes:
        atom_name=shape.name.split('_PAINT_',1)[0]
        shapes_by_atom.setdefault(atom_name,[]).append(shape)
    grouped=0;icon_grouped=0;component_grouped=0
    # Work from inner components out. An outer group can contain both existing
    # subgroups and leaf shapes without reparenting an already moved leaf twice.
    elements=list(ET.parse(master).getroot().iter())
    for element in reversed(elements):
        key=element.get('data-component-key') or element.get('data-icon-key')
        if not key: continue
        names={source_atoms[item.get('id')] for item in element.iter() if item.get('id') in source_atoms}
        unique={shape._element:shape for name in names for shape in shapes_by_atom.get(name,[])}
        order={node:i for i,node in enumerate(slide.shapes._spTree)}
        shapes=sorted(unique.values(),key=lambda shape:order[shape._element])
        if len(shapes)<2: continue
        first=order[shapes[0]._element]
        group=slide.shapes.add_group_shape(shapes)
        group.name='Figure component: '+key
        slide.shapes._spTree.remove(group._element)
        slide.shapes._spTree.insert(first,group._element)
        for name in names:shapes_by_atom[name]=[group]
        grouped+=1
        if element.get('data-component-key'):component_grouped+=1
        else:icon_grouped+=1
    def leaves(shapes):
        for shape in shapes:
            if shape.shape_type==6:yield from leaves(shape.shapes)
            else:yield shape
    leaf_shapes=list(leaves(slide.shapes))
    atom_sources={atom['objectName']:atom['sourceId'] for atom in cache['atoms']}
    for index,shape in enumerate(leaf_shapes):
        atom_name=shape.name.split('_PAINT_',1)[0]
        shape.name='fig_cos / '+atom_sources.get(atom_name,str(index))+f' / {index}'
    deck.save(str(pptx))
    picture_count = sum(1 for shape in leaf_shapes if shape.shape_type == 13)
    if picture_count:
        raise ValueError('Unexpected raster pictures in native export.')
    labels = [shape.text for shape in leaf_shapes if shape.has_text_frame and shape.text]
    if manifest:
        data = json.loads(manifest.read_text(encoding='utf-8-sig'))
        expected = [line for item in data['text_elements'] for line in str(item['content']).split('\n')]
        if Counter(expected) != Counter(labels):
            raise ValueError('Native text differs from the manifest; inspect output before delivery.')
    preview = output / 'svg-preview.png'
    try:
        import cairosvg
        cairosvg.svg2png(url=str(master), write_to=str(preview), output_width=1400)
    except (ImportError, OSError) as exc:
        warnings.append(f'SVG preview unavailable: {exc}')
    report = dict(ok=True, backend='offline-vtracer+native-ooxml' if args.input_image else 'offline-native-ooxml',
        source=str(input_path), source_sha256=source_hash, network_requests=0,
        external_service_credits=0, preset=args.preset if args.input_image else None,
        svg=str(master), pptx=str(pptx), svg_preview=str(preview) if preview.is_file() else None,
        ppt_render_verified=False, native_object_count=export['native_object_count'],
        native_text_count=len(labels), raster_picture_count=picture_count,
        icon_group_count=icon_grouped,component_group_count=component_grouped,native_group_count=grouped,
        top_level_object_count=len(slide.shapes),warnings=warnings,
        implicit_fill_closed_atoms=fill_closure_count,
        native_gradient_shapes=native_gradient_count,
        versions={name: importlib.metadata.version(name) for name in ['vtracer', 'python-pptx', 'fonttools']})
    (output / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f'FIG_DRAW_EXPORT_FAILED: {exc}', file=sys.stderr)
        raise SystemExit(1)
